from pathlib import Path

from openpyxl import load_workbook
import xlrd

from .models import BaseProduct
from .normalize import clean_text, normalize_code

CODE_NAMES = ("codigo", "código", "cod.", "cod ", "material", "referencia", "referência", "produto")
DESC_NAMES = ("descricao", "descrição", "desc.", "nome")


class SpreadsheetError(ValueError):
    pass


def _rows(path: Path) -> list[list[object]]:
    if path.suffix.lower() == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows: list[list[object]] = []
            for sheet in book.worksheets:
                rows.extend([list(row) for row in sheet.iter_rows(values_only=True)])
            return rows
        finally:
            book.close()
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(path)
        return [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for sheet in book.sheets() for r in range(sheet.nrows)]
    raise SpreadsheetError("Formato inválido. Envie um arquivo .xlsx ou .xls.")


def parse_spreadsheet(path: Path) -> list[BaseProduct]:
    rows = _rows(path)
    header_idx = code_col = None
    desc_col = None
    for row_idx, row in enumerate(rows[:50]):
        normalized = [clean_text(v).lower() for v in row]
        for col_idx, value in enumerate(normalized):
            if any(name in value for name in CODE_NAMES):
                header_idx, code_col = row_idx, col_idx
                break
        if code_col is not None:
            desc_col = next((i for i, value in enumerate(normalized) if any(name in value for name in DESC_NAMES)), None)
            break
    if code_col is None or header_idx is None:
        raise SpreadsheetError("Não foi possível identificar a coluna de código na planilha.")

    products: list[BaseProduct] = []
    seen: set[str] = set()
    for row in rows[header_idx + 1 :]:
        raw = row[code_col] if code_col < len(row) else None
        code, normalized = clean_text(raw), normalize_code(raw)
        if not normalized or normalized in seen:
            continue
        description = clean_text(row[desc_col]) if desc_col is not None and desc_col < len(row) else ""
        products.append(BaseProduct(code, normalized, description))
        seen.add(normalized)
    if not products:
        raise SpreadsheetError("A coluna identificada não contém códigos de produtos.")
    return products

